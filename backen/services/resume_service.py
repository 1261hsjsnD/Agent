import io
import os

from dotenv import load_dotenv

from agent.schemas.resume import ResumeProfile
from models.my_model import MyModel
# 再导出：简历上传接口直接从本模块导入 save_resume_to_student
from services.profile_service import save_resume_to_student  # noqa: F401

load_dotenv()

"""
简历服务
1. 模板管理：扫描 storage/resumes/ 目录，把成对的模板文件返回给前端
2. 简历解析：文件字节 -> 纯文本 -> 大模型 -> 结构化画像
"""

# 模板目录：storage/resumes/
_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # backen/
TEMPLATE_DIR = os.path.join(_BASE_DIR, "storage", "resumes")
TEMPLATE_DIR_NAME = "templates"

# 模板下载 URL 前缀（与 main.py 里挂载静态目录的前缀 /resume/templates 对应）
TEMPLATE_URL_PREFIX = "/resume/templates"

# 可下载的文档格式（默认取第一个，作为下载文件）
_DOWNLOAD_EXT = (".docx", ".doc")
# 可预览的格式（浏览器原生支持）
_PREVIEW_EXT = ".pdf"

# 文件名（不含扩展名）-> 展示名称
_LABEL_MAP = {
    "resume_template_java": "Java 工程师简历模板",
    "resume_template_marketing": "市场营销（国企）简历模板",
}


def _to_label(stem: str) -> str:
    """把文件名（不含扩展名）转成前端展示名称"""
    if stem in _LABEL_MAP:
        return _LABEL_MAP[stem]
    return stem.replace("_", " ")


def list_templates() -> list:
    """扫描模板目录，把同名文件配对成模板条目

    每个元素：
      {
        "label": 展示名,
        "download_url": 下载地址（docx）,
        "download_name": 下载时的文件名（带扩展名）,
        "preview_url": 预览地址（pdf，可能为空）,
      }
    """
    if not os.path.isdir(TEMPLATE_DIR):
        return []

    # 按「文件名（不含扩展名）」分组
    groups: dict = {}
    for filename in sorted(os.listdir(TEMPLATE_DIR)):
        full_path = os.path.join(TEMPLATE_DIR, filename)
        if not os.path.isfile(full_path):
            continue
        stem, ext = os.path.splitext(filename)
        ext = ext.lower()
        if ext not in _DOWNLOAD_EXT and ext != _PREVIEW_EXT:
            continue
        groups.setdefault(stem, {})[ext] = filename

    result = []
    for stem, files in groups.items():
        download_file = None
        for ext in _DOWNLOAD_EXT:
            if ext in files:
                download_file = files[ext]
                break
        # 没有任何可下载文档的模板跳过
        if not download_file:
            continue

        preview_file = files.get(_PREVIEW_EXT)
        label = _to_label(stem)
        download_ext = os.path.splitext(download_file)[1]

        result.append({
            "label": label,
            "download_url": f"{TEMPLATE_URL_PREFIX}/{download_file}",
            "download_name": f"{label}{download_ext}",
            "preview_url": f"{TEMPLATE_URL_PREFIX}/{preview_file}" if preview_file else "",
        })
    return result


# 简历解析提示词
RESUME_PROMPT = """
你是一名专业的简历解析助手。请从以下简历文本中提取信息，生成结构化的学生能力画像。

要求：
1. 专业技能（skills）：提取所有技术栈、编程语言、工具，并评估熟练度
2. 证书（certificates）：提取所有证书、资质
3. 项目经验（project_experience）：提取项目名称、技术栈、职责、成果，合并为一段描述
4. 创新能力（innovation）：从竞赛、专利、创新项目中提炼
5. 学习能力（learning_ability）：从自学经历、绩点、技术学习中提炼
6. 抗压能力（pressure_ability）：从高强度经历、实习表现中提炼
7. 沟通能力（communication）：从社团、演讲、团队协作中提炼
8. 实习经历（internships）：提取实习公司、岗位、时长、内容
如果某项信息缺失，对应字段留空即可。
"""


def analysis_resume(content: bytes, filename: str) -> ResumeProfile:
    """完整解析：文件字节 -> 纯文本 -> 在线大模型 -> 结构化画像"""
    # 1. 提取纯文本
    text = _extract_text(content, filename)

    # 2. 截断超长文本，防止超 token 限制
    if len(text) > 6000:
        text = text[:6000]

    # 3. 用在线模型 + 结构化输出解析画像
    model = MyModel.get_model()
    structured_model = model.with_structured_output(ResumeProfile)

    result: ResumeProfile = structured_model.invoke([
        ("system", RESUME_PROMPT),
        ("human", text),
    ])

    return result


def _extract_text(content: bytes, filename: str) -> str:
    """根据文件类型提取纯文本"""
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return _parse_pdf(content)
    elif ext == "docx":
        return _parse_docx(content)
    else:
        raise ValueError(f"不支持的简历格式: {ext}")


def _parse_pdf(content: bytes) -> str:
    from pdfminer.high_level import extract_text
    text = extract_text(io.BytesIO(content))
    return text


def _parse_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    data = []
    # 先读正文段落
    for p in doc.paragraphs:
        if p.text.strip():
            data.append(p.text)
    # 再读表格
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            data.append(" | ".join(row_data))
    return "\n".join(data)
